#!/usr/bin/env python3
"""
将 ReadCode/ 目录下的 Markdown 分析文章转换为 PowerPoint 演示文稿。
每个 MD 文件生成一个对应的 PPTX 文件，保存到 ReadCode/PPT/ 目录。
"""

import os
import re
import glob
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ─── 配色方案 ───
COLOR_TITLE_BG = RGBColor(0x1B, 0x2A, 0x4A)      # 深蓝标题栏背景
COLOR_TITLE_TEXT = RGBColor(0xFF, 0xFF, 0xFF)      # 白色标题文字
COLOR_SUBTITLE = RGBColor(0xCC, 0xCC, 0xCC)        # 灰色副标题
COLOR_BODY_TEXT = RGBColor(0x33, 0x33, 0x33)        # 深灰正文
COLOR_ACCENT = RGBColor(0x2E, 0x86, 0xC1)          # 蓝色强调
COLOR_CODE_BG = RGBColor(0xF5, 0xF5, 0xF5)         # 浅灰代码背景
COLOR_CODE_TEXT = RGBColor(0x2C, 0x3E, 0x50)        # 深灰代码文字
COLOR_MERMAID_BG = RGBColor(0xEB, 0xF5, 0xFB)      # 浅蓝 Mermaid 背景
COLOR_MERMAID_TEXT = RGBColor(0x1A, 0x5C, 0x8A)    # 蓝色 Mermaid 文字
COLOR_LINK = RGBColor(0x2E, 0x86, 0xC1)            # 蓝色链接
COLOR_SECTION_BG = RGBColor(0x1B, 0x2A, 0x4A)      # 章节分隔页背景
COLOR_SECTION_TEXT = RGBColor(0xFF, 0xFF, 0xFF)     # 章节分隔页文字
COLOR_COVER_BG = RGBColor(0x0D, 0x1B, 0x2A)        # 封面背景
COLOR_COVER_ACCENT = RGBColor(0x2E, 0x86, 0xC1)    # 封面强调线
COLOR_BULLET = RGBColor(0x2E, 0x86, 0xC1)          # 项目符号颜色
COLOR_TABLE_HEADER = RGBColor(0x1B, 0x2A, 0x4A)    # 表头背景
COLOR_TABLE_ALT = RGBColor(0xF0, 0xF4, 0xF8)       # 表格交替行

# ─── 尺寸常量 ───
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)
TITLE_BAR_HEIGHT = Inches(1.0)
CONTENT_LEFT = Inches(0.8)
CONTENT_TOP = Inches(1.3)
CONTENT_WIDTH = Inches(11.733)
CONTENT_HEIGHT = Inches(5.8)
FONT_TITLE = "Microsoft YaHei"
FONT_BODY = "Microsoft YaHei"
FONT_CODE = "Consolas"


def create_presentation():
    """创建 16:9 宽屏演示文稿"""
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    return prs


def add_shape(slide, left, top, width, height, fill_color=None):
    """添加矩形形状"""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.line.fill.background()
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    return shape


def add_textbox(slide, left, top, width, height, text="", font_size=14,
                font_color=COLOR_BODY_TEXT, font_name=FONT_BODY, bold=False,
                alignment=PP_ALIGN.LEFT, word_wrap=True):
    """添加文本框"""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = word_wrap
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = font_color
    p.font.name = font_name
    p.font.bold = bold
    p.alignment = alignment
    return txBox


def add_cover_slide(prs, title, subtitle="Headroom 项目代码深度分析"):
    """添加封面页"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白布局

    # 背景
    bg = add_shape(slide, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT, COLOR_COVER_BG)

    # 强调线
    add_shape(slide, Inches(1.5), Inches(3.0), Inches(2.0), Pt(4), COLOR_COVER_ACCENT)

    # 标题
    add_textbox(slide, Inches(1.5), Inches(3.2), Inches(10.0), Inches(1.5),
                text=title, font_size=36, font_color=COLOR_TITLE_TEXT,
                font_name=FONT_TITLE, bold=True)

    # 副标题
    add_textbox(slide, Inches(1.5), Inches(4.7), Inches(10.0), Inches(0.8),
                text=subtitle, font_size=18, font_color=COLOR_SUBTITLE,
                font_name=FONT_BODY)

    return slide


def add_section_slide(prs, section_title):
    """添加章节分隔页（总/分/总）"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # 全屏背景
    add_shape(slide, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT, COLOR_SECTION_BG)

    # 强调线
    add_shape(slide, Inches(2.0), Inches(3.2), Inches(1.5), Pt(3), COLOR_COVER_ACCENT)

    # 章节标题
    add_textbox(slide, Inches(2.0), Inches(3.5), Inches(9.0), Inches(1.2),
                text=section_title, font_size=32, font_color=COLOR_SECTION_TEXT,
                font_name=FONT_TITLE, bold=True, alignment=PP_ALIGN.LEFT)

    return slide


def add_content_slide(prs, title, content_lines):
    """添加内容幻灯片"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # 标题栏背景
    add_shape(slide, 0, 0, SLIDE_WIDTH, TITLE_BAR_HEIGHT, COLOR_TITLE_BG)

    # 标题文字
    add_textbox(slide, Inches(0.8), Inches(0.15), Inches(11.5), Inches(0.7),
                text=title, font_size=22, font_color=COLOR_TITLE_TEXT,
                font_name=FONT_TITLE, bold=True)

    # 内容区域
    if not content_lines:
        return slide

    txBox = slide.shapes.add_textbox(CONTENT_LEFT, CONTENT_TOP, CONTENT_WIDTH, CONTENT_HEIGHT)
    tf = txBox.text_frame
    tf.word_wrap = True

    for i, (text, style) in enumerate(content_lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()

        p.text = text
        p.font.name = style.get("font", FONT_BODY)
        p.font.size = Pt(style.get("size", 14))
        p.font.color.rgb = style.get("color", COLOR_BODY_TEXT)
        p.font.bold = style.get("bold", False)
        p.alignment = style.get("align", PP_ALIGN.LEFT)
        p.space_after = Pt(style.get("space_after", 4))
        p.space_before = Pt(style.get("space_before", 0))

        if style.get("indent"):
            p.level = style["indent"]

    return slide


def add_mermaid_slide(prs, chart_title, mermaid_code):
    """添加 Mermaid 图表幻灯片"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # 标题栏
    add_shape(slide, 0, 0, SLIDE_WIDTH, TITLE_BAR_HEIGHT, COLOR_TITLE_BG)
    add_textbox(slide, Inches(0.8), Inches(0.15), Inches(11.5), Inches(0.7),
                text=chart_title, font_size=22, font_color=COLOR_TITLE_TEXT,
                font_name=FONT_TITLE, bold=True)

    # Mermaid 背景
    add_shape(slide, Inches(0.5), Inches(1.2), Inches(12.333), Inches(5.8), COLOR_MERMAID_BG)

    # Mermaid 标签
    add_textbox(slide, Inches(0.8), Inches(1.3), Inches(3.0), Inches(0.4),
                text="Mermaid 图表", font_size=11, font_color=COLOR_MERMAID_TEXT,
                font_name=FONT_BODY, bold=True)

    # Mermaid 代码
    # 截断过长的代码
    lines = mermaid_code.strip().split("\n")
    display_code = "\n".join(lines[:40])
    if len(lines) > 40:
        display_code += f"\n... (共 {len(lines)} 行)"

    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.733), Inches(5.0))
    tf = txBox.text_frame
    tf.word_wrap = True

    for i, line in enumerate(display_code.split("\n")):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = line
        p.font.name = FONT_CODE
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_MERMAID_TEXT
        p.space_after = Pt(1)
        p.space_before = Pt(0)

    return slide


def add_code_slide(prs, title, code_text):
    """添加代码幻灯片"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # 标题栏
    add_shape(slide, 0, 0, SLIDE_WIDTH, TITLE_BAR_HEIGHT, COLOR_TITLE_BG)
    add_textbox(slide, Inches(0.8), Inches(0.15), Inches(11.5), Inches(0.7),
                text=title, font_size=22, font_color=COLOR_TITLE_TEXT,
                font_name=FONT_TITLE, bold=True)

    # 代码背景
    add_shape(slide, Inches(0.5), Inches(1.2), Inches(12.333), Inches(5.8), COLOR_CODE_BG)

    # 代码文本
    lines = code_text.strip().split("\n")
    display_code = "\n".join(lines[:35])
    if len(lines) > 35:
        display_code += f"\n... (共 {len(lines)} 行)"

    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.4))
    tf = txBox.text_frame
    tf.word_wrap = True

    for i, line in enumerate(display_code.split("\n")):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = line
        p.font.name = FONT_CODE
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_CODE_TEXT
        p.space_after = Pt(1)
        p.space_before = Pt(0)

    return slide


def add_table_slide(prs, title, headers, rows):
    """添加表格幻灯片"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # 标题栏
    add_shape(slide, 0, 0, SLIDE_WIDTH, TITLE_BAR_HEIGHT, COLOR_TITLE_BG)
    add_textbox(slide, Inches(0.8), Inches(0.15), Inches(11.5), Inches(0.7),
                text=title, font_size=22, font_color=COLOR_TITLE_TEXT,
                font_name=FONT_TITLE, bold=True)

    if not headers or not rows:
        return slide

    # 表格
    num_rows = len(rows) + 1
    num_cols = len(headers)
    table_width = min(Inches(11.733), Inches(num_cols * 2.5))
    table_left = Inches(0.8)

    table_shape = slide.shapes.add_table(num_rows, num_cols, table_left, Inches(1.5),
                                          table_width, Inches(0.5 * num_rows))
    table = table_shape.table

    # 表头
    for j, header in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = header
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.size = Pt(12)
            paragraph.font.bold = True
            paragraph.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            paragraph.font.name = FONT_BODY
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_TABLE_HEADER

    # 数据行
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            if j < num_cols:
                cell = table.cell(i + 1, j)
                cell.text = str(val)[:100]  # 截断过长文本
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.font.size = Pt(10)
                    paragraph.font.color.rgb = COLOR_BODY_TEXT
                    paragraph.font.name = FONT_BODY
                if i % 2 == 1:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = COLOR_TABLE_ALT

    return slide


def parse_markdown(md_text):
    """解析 Markdown 文本为结构化数据"""
    lines = md_text.split("\n")
    sections = []
    current_section = None
    current_content = []
    in_code_block = False
    code_block = []
    code_lang = ""
    in_mermaid = False
    mermaid_block = []
    mermaid_title = ""

    i = 0
    while i < len(lines):
        line = lines[i]

        # Mermaid 代码块
        if line.strip().startswith("```mermaid"):
            in_mermaid = True
            mermaid_block = []
            # 尝试从前面的 subgraph 或注释中提取标题
            mermaid_title = ""
            i += 1
            continue

        if in_mermaid:
            if line.strip() == "```":
                in_mermaid = False
                if current_section:
                    current_content.append(("mermaid", "\n".join(mermaid_block), mermaid_title))
                i += 1
                continue
            # 提取 subgraph 标题
            if "subgraph" in line and not mermaid_title:
                match = re.search(r'subgraph\s+(.+)', line)
                if match:
                    mermaid_title = match.group(1).strip()
            mermaid_block.append(line)
            i += 1
            continue

        # 普通代码块
        if line.strip().startswith("```") and not in_code_block:
            in_code_block = True
            code_block = []
            code_lang = line.strip()[3:].strip()
            i += 1
            continue

        if in_code_block:
            if line.strip() == "```":
                in_code_block = False
                if current_section:
                    current_content.append(("code", "\n".join(code_block), code_lang))
                i += 1
                continue
            code_block.append(line)
            i += 1
            continue

        # 标题
        if line.startswith("# ") and not line.startswith("## "):
            # 一级标题 → 封面
            if current_section:
                sections.append((current_section, current_content))
            current_section = ("cover", line[2:].strip())
            current_content = []
            i += 1
            continue

        if line.startswith("## "):
            if current_section:
                sections.append((current_section, current_content))
            current_section = ("h2", line[3:].strip())
            current_content = []
            i += 1
            continue

        if line.startswith("### "):
            if current_section:
                sections.append((current_section, current_content))
            current_section = ("h3", line[4:].strip())
            current_content = []
            i += 1
            continue

        if line.startswith("#### "):
            # 四级标题作为内容的一部分
            current_content.append(("h4", line[5:].strip(), ""))
            i += 1
            continue

        # 表格行
        if "|" in line and line.strip().startswith("|"):
            table_lines = []
            while i < len(lines) and "|" in lines[i] and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            # 解析表格
            parsed = parse_table(table_lines)
            if parsed:
                current_content.append(("table", parsed, ""))
            continue

        # 列表项
        if re.match(r'^\s*[-*]\s', line):
            text = re.sub(r'^\s*[-*]\s+', '', line).strip()
            # 清理 markdown 链接
            text = clean_md_text(text)
            current_content.append(("bullet", text, ""))
            i += 1
            continue

        # 有序列表
        if re.match(r'^\s*\d+\.\s', line):
            text = re.sub(r'^\s*\d+\.\s+', '', line).strip()
            text = clean_md_text(text)
            current_content.append(("numbered", text, ""))
            i += 1
            continue

        # 普通段落
        stripped = line.strip()
        if stripped:
            text = clean_md_text(stripped)
            current_content.append(("text", text, ""))
        i += 1

    if current_section:
        sections.append((current_section, current_content))

    return sections


def parse_table(table_lines):
    """解析 Markdown 表格"""
    if len(table_lines) < 2:
        return None

    def parse_row(line):
        cells = [c.strip() for c in line.strip("|").split("|")]
        return [clean_md_text(c) for c in cells]

    headers = parse_row(table_lines[0])

    # 跳过分隔行
    data_start = 1
    if len(table_lines) > 1 and re.match(r'^[\s|:-]+$', table_lines[1]):
        data_start = 2

    rows = []
    for line in table_lines[data_start:]:
        row = parse_row(line)
        if row:
            rows.append(row)

    return (headers, rows)


def clean_md_text(text):
    """清理 Markdown 格式文本"""
    # 保留链接文本但去掉 URL
    text = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', text)
    # 去掉加粗标记
    text = re.sub(r'\*\*([^*]+)\*\*', r'\1', text)
    # 去掉斜体标记
    text = re.sub(r'\*([^*]+)\*', r'\1', text)
    # 去掉行内代码标记
    text = re.sub(r'`([^`]+)`', r'\1', text)
    # 去掉 HTML 标签
    text = re.sub(r'<[^>]+>', '', text)
    return text.strip()


def md_to_pptx(md_path, output_path):
    """将 Markdown 文件转换为 PPTX"""
    with open(md_path, "r", encoding="utf-8") as f:
        md_text = f.read()

    sections = parse_markdown(md_text)
    prs = create_presentation()

    for section_info, content in sections:
        level, title = section_info

        # 封面页
        if level == "cover":
            add_cover_slide(prs, title)
            continue

        # 章节分隔页（## 总 / ## 分 / ## 总）
        if level == "h2":
            # 判断是否为"总—分—总"的章节分隔
            if any(kw in title for kw in ["总（", "分（", "总（"]):
                add_section_slide(prs, title)
                continue

            # 普通 h2 作为内容页
            content_lines = build_content_lines(content)
            if content_lines:
                add_content_slide(prs, title, content_lines)
            else:
                add_section_slide(prs, title)
            continue

        # h3 内容页
        if level == "h3":
            # 先处理 Mermaid 和大代码块为独立幻灯片
            remaining = []
            for item in content:
                item_type, item_data, item_meta = item
                if item_type == "mermaid":
                    chart_title = item_meta if item_meta else title + " — 图表"
                    add_mermaid_slide(prs, chart_title, item_data)
                elif item_type == "code" and len(item_data.split("\n")) > 10:
                    code_title = title + (" — " + item_meta if item_meta else " — 代码")
                    add_code_slide(prs, code_title, item_data)
                else:
                    remaining.append(item)

            if remaining:
                content_lines = build_content_lines(remaining)
                add_content_slide(prs, title, content_lines)
            continue

    prs.save(output_path)
    return output_path


def build_content_lines(content_items):
    """将内容项转换为幻灯片文本行"""
    lines = []
    for item in content_items:
        item_type, item_data, item_meta = item

        if item_type == "bullet":
            lines.append((f"• {item_data}", {
                "size": 13, "color": COLOR_BODY_TEXT, "space_after": 3
            }))

        elif item_type == "numbered":
            lines.append((f"  {item_data}", {
                "size": 13, "color": COLOR_BODY_TEXT, "space_after": 3
            }))

        elif item_type == "text":
            lines.append((item_data, {
                "size": 13, "color": COLOR_BODY_TEXT, "space_after": 4
            }))

        elif item_type == "h4":
            lines.append((item_data, {
                "size": 15, "color": COLOR_ACCENT, "bold": True, "space_before": 8, "space_after": 3
            }))

        elif item_type == "code":
            # 短代码块内嵌
            code_lines = item_data.strip().split("\n")
            display = "\n".join(code_lines[:8])
            if len(code_lines) > 8:
                display += " ..."
            lang_label = f" [{item_meta}]" if item_meta else ""
            lines.append((f"代码{lang_label}:", {
                "size": 10, "color": COLOR_ACCENT, "bold": True, "space_before": 6, "space_after": 2
            }))
            lines.append((display, {
                "size": 9, "color": COLOR_CODE_TEXT, "font": FONT_CODE, "space_after": 4
            }))

        elif item_type == "table":
            # 表格在内容行中以简略形式展示
            headers, rows = item_data
            lines.append((f"表格: {' | '.join(headers[:5])}", {
                "size": 11, "color": COLOR_ACCENT, "bold": True, "space_before": 6, "space_after": 2
            }))
            for row in rows[:5]:
                lines.append(("  " + " | ".join(str(c)[:30] for c in row[:5]), {
                    "size": 10, "color": COLOR_BODY_TEXT, "space_after": 1
                }))
            if len(rows) > 5:
                lines.append((f"  ... 共 {len(rows)} 行", {
                    "size": 10, "color": COLOR_SUBTITLE, "space_after": 2
                }))

    return lines


def main():
    md_dir = "/workspace/ReadCode"
    ppt_dir = "/workspace/ReadCode/PPT"
    os.makedirs(ppt_dir, exist_ok=True)

    md_files = sorted(glob.glob(os.path.join(md_dir, "*.md")))
    print(f"找到 {len(md_files)} 个 Markdown 文件")

    success = 0
    failed = 0

    for md_path in md_files:
        basename = os.path.basename(md_path)
        name_without_ext = os.path.splitext(basename)[0]
        pptx_path = os.path.join(ppt_dir, f"{name_without_ext}.pptx")

        try:
            md_to_pptx(md_path, pptx_path)
            print(f"  ✓ {basename} → {name_without_ext}.pptx")
            success += 1
        except Exception as e:
            print(f"  ✗ {basename} → 失败: {e}")
            failed += 1

    print(f"\n完成：成功 {success} 个，失败 {failed} 个")
    print(f"PPT 文件保存在：{ppt_dir}")


if __name__ == "__main__":
    main()
